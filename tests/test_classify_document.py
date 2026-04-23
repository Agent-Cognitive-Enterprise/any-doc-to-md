"""
Tests for anydoc2md.format_converters.classification.classify_document.

These tests use synthetic files created inside tmp_path so the package has
no dependency on a parent repository test corpus.
"""
from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest

from anydoc2md.format_converters.classification.classify_document import (
    DocumentTraits,
    PDF_TABLE_SCAN_MAX_PAGES,
    classify,
    _classify_pdf,
    _unknown_traits,
)


class TestDocumentTraits:
    def test_to_dict_has_all_fields(self) -> None:
        t = _unknown_traits("pdf")
        d = t.to_dict()
        for key in (
            "file_type",
            "page_count",
            "image_count",
            "table_count",
            "word_count",
            "is_scanned",
            "is_image_heavy",
            "is_table_heavy",
            "is_multi_column",
            "is_text_only",
            "has_math",
        ):
            assert key in d

    def test_unknown_traits_defaults(self) -> None:
        t = _unknown_traits("pdf")
        assert t.file_type == "pdf"
        assert t.is_text_only is True
        assert t.image_count == 0


class TestDispatcher:
    def test_unsupported_extension_returns_unknown(self, tmp_path: Path) -> None:
        f = tmp_path / "doc.xyz"
        f.write_text("data", encoding="utf-8")
        t = classify(f)
        assert t.file_type == "xyz"

    def test_missing_file_does_not_raise(self, tmp_path: Path) -> None:
        t = classify(tmp_path / "missing.pdf")
        assert t is not None


class TestPdfTableScanPolicy:
    def test_large_pdf_table_scan_is_capped_and_suppresses_layout_recommendation(
        self,
        tmp_path: Path,
    ) -> None:
        page_count = PDF_TABLE_SCAN_MAX_PAGES + 10
        pages = [_FakePdfPage() for _ in range(page_count)]
        fake_doc = _FakePdfDoc(pages)
        fake_fitz = SimpleNamespace(
            open=MagicMock(return_value=fake_doc),
            no_recommend_layout=MagicMock(),
        )

        with patch.dict("sys.modules", {"fitz": fake_fitz}):
            traits = _classify_pdf(tmp_path / "large.pdf")

        assert traits.page_count == page_count
        assert fake_fitz.no_recommend_layout.call_count == 1
        assert sum(page.find_tables_call_count for page in pages) == PDF_TABLE_SCAN_MAX_PAGES
        assert pages[0].find_tables_call_count == 1
        assert pages[-1].find_tables_call_count == 1

    def test_table_scan_stops_after_table_heavy_threshold(self, tmp_path: Path) -> None:
        pages = [_FakePdfPage(table_count=2)] + [_FakePdfPage() for _ in range(5)]
        fake_doc = _FakePdfDoc(pages)
        fake_fitz = SimpleNamespace(
            open=MagicMock(return_value=fake_doc),
            no_recommend_layout=MagicMock(),
        )

        with patch.dict("sys.modules", {"fitz": fake_fitz}):
            traits = _classify_pdf(tmp_path / "table-heavy.pdf")

        assert traits.table_count == 2
        assert traits.is_table_heavy is True
        assert sum(page.find_tables_call_count for page in pages) == 1


class TestSyntheticTextAndHtml:
    def test_txt_detects_math(self, tmp_path: Path) -> None:
        f = tmp_path / "math.txt"
        f.write_text("The formula is $E = mc^2$ and \\frac{a}{b}.", encoding="utf-8")
        t = classify(f)
        assert t.file_type == "txt"
        assert t.has_math is True

    def test_html_detects_images_and_tables(self, tmp_path: Path) -> None:
        f = tmp_path / "doc.html"
        f.write_text(
            "<html><body>"
            + '<img src="a.png">' * 5
            + "<table><tr><td>x</td></tr></table>" * 2
            + "</body></html>",
            encoding="utf-8",
        )
        t = classify(f)
        assert t.file_type == "html"
        assert t.image_count == 5
        assert t.table_count == 2
        assert t.is_image_heavy is True
        assert t.is_table_heavy is True


@pytest.mark.skipif(pytest.importorskip("fitz", reason="PyMuPDF required") is None, reason="PyMuPDF required")
class TestSyntheticPdf:
    def test_text_only_pdf_classifies_as_pdf(self, tmp_path: Path) -> None:
        pdf = tmp_path / "text-only.pdf"
        _write_text_pdf(pdf)
        t = classify(pdf)
        assert t.file_type == "pdf"
        assert t.page_count == 1
        assert t.word_count > 0
        assert t.is_text_only is True
        assert t.image_count == 0

    def test_multi_column_pdf_detected(self, tmp_path: Path) -> None:
        pdf = tmp_path / "multi-column.pdf"
        _write_multi_column_pdf(pdf)
        t = classify(pdf)
        assert t.file_type == "pdf"
        assert t.is_multi_column is True

    def test_image_pdf_not_text_only(self, tmp_path: Path) -> None:
        pdf = tmp_path / "with-image.pdf"
        _write_image_pdf(pdf)
        t = classify(pdf)
        assert t.file_type == "pdf"
        assert t.image_count >= 1
        assert t.is_text_only is False

    def test_multiple_generated_pdfs_finish_and_have_nonnegative_counts(self, tmp_path: Path) -> None:
        pdfs = [
            tmp_path / "one.pdf",
            tmp_path / "two.pdf",
            tmp_path / "three.pdf",
        ]
        _write_text_pdf(pdfs[0])
        _write_multi_column_pdf(pdfs[1])
        _write_image_pdf(pdfs[2])

        for pdf in pdfs:
            t = classify(pdf)
            assert t.file_type == "pdf"
            assert t.page_count > 0
            assert t.image_count >= 0
            assert t.table_count >= 0
            assert t.word_count >= 0


@pytest.mark.skipif(pytest.importorskip("docx", reason="python-docx required") is None, reason="python-docx required")
class TestSyntheticDocx:
    def test_docx_with_table_classifies(self, tmp_path: Path) -> None:
        docx_path = tmp_path / "table.docx"
        _write_docx_with_table(docx_path)
        t = classify(docx_path)
        assert t.file_type == "docx"
        assert t.table_count >= 1


def _write_text_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello world " * 80)
    doc.save(path)
    doc.close()


def _write_multi_column_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    for idx in range(3):
        top = 48 + (idx * 180)
        left = fitz.Rect(40, top, 220, top + 120)
        right = fitz.Rect(380, top, 555, top + 120)
        left_text = ("Left column block %d. " % idx) * 10
        right_text = ("Right column block %d. " % idx) * 10
        page.insert_textbox(left, left_text, fontsize=11)
        page.insert_textbox(right, right_text, fontsize=11)
    doc.save(path)
    doc.close()


def _write_image_pdf(path: Path) -> None:
    import fitz
    from PIL import Image

    img = Image.new("RGB", (12, 12), color=(255, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    image_bytes = buf.getvalue()

    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(72, 72, 144, 144), stream=image_bytes)
    doc.save(path)
    doc.close()


def _write_docx_with_table(path: Path) -> None:
    import docx

    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    doc.save(path)


class _FakePdfPage:
    def __init__(self, *, table_count: int = 0) -> None:
        self._table_count = table_count
        self.find_tables_call_count = 0
        self.rect = SimpleNamespace(width=612)

    def get_images(self, *, full: bool) -> list:
        return []

    def get_text(self, kind: str, **kwargs) -> str | list:
        if kind == "blocks":
            return []
        return "hello world sample document text"

    def find_tables(self):
        self.find_tables_call_count += 1
        return SimpleNamespace(tables=[object()] * self._table_count)


class _FakePdfDoc:
    def __init__(self, pages: list[_FakePdfPage]) -> None:
        self._pages = pages
        self.closed = False

    def __len__(self) -> int:
        return len(self._pages)

    def __iter__(self):
        return iter(self._pages)

    def extract_image(self, xref: int) -> dict:
        return {"image": b""}

    def close(self) -> None:
        self.closed = True
